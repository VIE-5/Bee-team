# -*- coding: utf-8 -*-
"""
标志方案展示 · 通用版式构建（母版 · 5 页）
=================================================================
版式来源：碗香 Logo「一」确立、经「二」校验的 5 页结构
   P1-P3  每版详页：左标志展示台 + 右「标志构成」三段释义 + 同页色彩应用
   P4     三版并置总览
   P5     三版综合应用（门头灯箱 / 外带碗盖 / 餐具烙印 / 员工围裙 / 名片）

画布 1360x850 (16:10)，1 px = 9525 EMU —— 全案母版几何 1:1 直落 PPT。
产物为**真文本 + 真图片 + 真色块**，客户可直接改字换图。

用法：改下面【配置区】（路径 / 标题 / VERSIONS 数据），然后
    python build.py
素材须先用 prep-assets.py 加工到 <WK>/assets/trim/。
=================================================================
"""
import os
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import Image, ImageFont

# ==================== 配置区（只改这一段） ====================
WK = r'D:\...\_work'                       # 工作目录（内含 assets/trim/）
OUT = r'D:\...\01-工蜂产出\标志方案展示.pptx'

BRAND_NAME = '碗香'                         # 页眉左侧品牌字
BRAND_EN = 'WAN XIANG'                     # 页眉右侧英文
BRAND = 'C63021'                           # 页面强调色（红短线 / 编号）
BRAND_D = '992D27'                         # 小标题色

# 各版数据。key 对应 prep-assets.py 的 FILES key（A/B/C，建议 3 版）
#   no              编号（壹/贰/叁），同时用于详情页页眉与并置页
#   blocks          详情页三段释义：(小标题, 正文) 各 60~80 字，控制在 3 行内
#   colors          色彩应用：(名称, HEX, 建议占比%, 卡面文字色 '1C1C16' | 'FFFFFF')
#                   占比合计须 = 100；某段宽 < 72px 时百分比标注自动略过
#   name / tags     并置页：版名、适用场景
VERSIONS = [
    dict(
        no='壹', key='A', name='红金复古国潮',
        blocks=[
            ('徽章结构', '暗红实心圆打底，外环与内圆之间以双道金线分界，左右各嵌一枚金色菱形作配重。整体走老字号正章的路子，边线利落，远看即见分量。'),
            ('人物主体', '半写实男厨师居中抱臂，白厨师服配黑围裙，短发利落、朗笑露齿，目光与观者平接。肤色与衣褶按插画方式分面处理，庄重而不失亲和。'),
            ('字标与标语', '顶部沿弧线排「乡土黔菜」四字，下方牌匾落「碗香」毛笔主标，配「小炒剔骨鸡」副标，左右金色横线收边，菜系与招牌一次交代完整。'),
        ],
        colors=[('暗红', '8A0F0D', 56, 'FFFFFF'), ('香槟金', 'F5D68F', 22, '1C1C16'),
                ('米粉白', 'FAEBD2', 14, '1C1C16'), ('墨黑', '130501', 8, 'FFFFFF')],
        note='暗红为主色，香槟金承担顶标、金线与菱形，米粉白为字标与制服，墨黑仅作围裙与暗部',
        tags='门头灯箱 / 主 KV / 堂食招牌',
    ),
    dict(
        no='贰', key='B', name='深绿米黄乡土',
        blocks=[
            ('徽章结构', '外环改墨绿实心，内圆填米黄并压一道同心细线，明暗关系比壹版更缓。菱形与顶标同取米黄，整体更像一方旧时印版。'),
            ('人物主体', '沿用同一厨师形象与姿态，改以米黄暖光衬底。白衣在深绿环中跳脱而出，人物轮廓比壹版更清晰，小尺寸下的辨识度最高。'),
            ('字标与标语', '顶标换作「鲜炒贵州味」，直说烹饪方式；牌匾同用「碗香」主标与「小炒剔骨鸡」副标，深绿底衬米黄字，是三版里对比最干净的一组。'),
        ],
        colors=[('墨绿', '074330', 42, 'FFFFFF'), ('米黄', 'FDE4AD', 34, '1C1C16'),
                ('纸白', 'FFFDF6', 14, '1C1C16'), ('深绿', '052415', 10, 'FFFFFF')],
        note='墨绿为主色（外环与牌匾），米黄承担内圆与字标，纸白为制服，深绿作暗部与细线',
        tags='外带包装 / 社媒头像 / 社区门店',
    ),
    dict(
        no='叁', key='C', name='朱红米黄',
        blocks=[
            ('徽章结构', '反过来以米黄作外环、朱红作内圆，形成「浅包深」的收束感。去掉全部金线，仅留一圈极细的深色描边，气质比前两版更素净。'),
            ('人物主体', '同一厨师、同一姿态。因内圆由暗色转为亮朱红，白衣与红底之间的反差最强，人物仿佛被一束顶光打亮，面部最显干净。'),
            ('字标与标语', '顶标回到「乡土黔菜」，改以墨色书写；牌匾主标提为「剔骨小炒鸡」，把招牌菜名放到第一层级，副标「贵州土菜」用于点明菜系。'),
        ],
        colors=[('朱红', 'A80F0D', 46, 'FFFFFF'), ('米黄', 'FAE1C0', 32, '1C1C16'),
                ('纸白', 'FFFCF5', 12, '1C1C16'), ('墨黑', '0B0706', 10, 'FFFFFF')],
        note='朱红为主色（内圆与牌匾），米黄承担外环与主标，纸白为制服，墨黑用于顶标、围裙与描边',
        tags='菜单封面 / 招牌菜物料 / 正式文件',
    ),
]

# P4 主标题；P4/P5 底部各一行说明；P1-P3 色彩应用右上角提示语
LEAD_P4_TITLE = '同一形象骨架 · 三套色彩方案'
LEAD_P4 = '三版共用同一形象与牌匾骨架，靠配色与顶标用字区分使用场合'
FOOT_P4 = '三版为同一品牌的三套应用形态，可同时使用，但同一物料中不得混排两版标志'
FOOT_P5 = '三版按场景分派：壹版压门头与主 KV，贰版用于外带与社媒，叁版用于菜单封面与正式物料'
COLOR_HINT = '色值实测取自本版标志 · 配比为建议应用占比'
# ============================================================

A = os.path.join(WK, 'assets', 'trim')
OUTDIR = os.path.dirname(OUT)
os.makedirs(OUTDIR, exist_ok=True)

PXU = 9525
W, H = 1360, 850
MX = W * 0.11          # 149.6  左右边距 11%
HDR = H * 0.15         # 127.5  页眉带 15%
CW = W - 2 * MX        # 1060.8 版心宽

INK = '1C1C16'; INK2 = '6E6A5C'
PAPER = 'FFFFFF'; PLATE = 'FAF7F2'; LINEC = 'C6C2B8'; DARK = '1D191A'

FCN = 'Microsoft YaHei'; FEN = 'Arial'; FMONO = 'Consolas'
FPATH = {
    (FCN, True): r'C:\Windows\Fonts\msyhbd.ttc',
    (FCN, False): r'C:\Windows\Fonts\msyh.ttc',
    (FEN, False): r'C:\Windows\Fonts\arial.ttf',
    (FEN, True): r'C:\Windows\Fonts\arialbd.ttf',
}
_fc = {}


def _font(fam, bold):
    k = (fam, bold)
    if k not in _fc:
        p = FPATH.get(k) or FPATH[(FCN, bold)]
        try:
            _fc[k] = ImageFont.truetype(p, 100)
        except OSError:
            _fc[k] = ImageFont.truetype(FPATH[(FCN, bold)], 100)
    return _fc[k]


def E(v):
    return Emu(int(round(v * PXU)))


def tw(s, px, bold=False, fam=FCN):
    return _font(fam, bold).getlength(s) * px / 100


# ---------- 基础件 ----------
def rect(slide, x, y, w, h, color, line=None, weight=1.0, radius=None):
    shp = (MSO_SHAPE.ROUNDED_RECTANGLE if radius is not None else MSO_SHAPE.RECTANGLE)
    s = slide.shapes.add_shape(shp, E(x), E(y), E(w), E(h))
    if radius is not None:
        s.adjustments[0] = radius
    s.fill.solid(); s.fill.fore_color.rgb = RGBColor.from_string(color)
    if line:
        s.line.color.rgb = RGBColor.from_string(line)
        s.line.width = Pt(weight * 0.75)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def frame(slide, x, y, w, h, line=LINEC, weight=1.0, fill=None):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, E(x), E(y), E(w), E(h))
    if fill:
        s.fill.solid(); s.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        s.fill.background()
    s.line.color.rgb = RGBColor.from_string(line)
    s.line.width = Pt(weight * 0.75)
    s.shadow.inherit = False
    return s


def oval(slide, x, y, w, h, fill=None, line=None, weight=1.0):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, E(x), E(y), E(w), E(h))
    if fill:
        s.fill.solid(); s.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = RGBColor.from_string(line)
        s.line.width = Pt(weight * 0.75)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


ALIGN = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT, 'j': PP_ALIGN.JUSTIFY}
ANCH = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}


def text(slide, x, y, w, h, s, size, color, bold=False, align='l', fam=FCN,
         spc=0.0, lh=None, anchor='t'):
    """size/lh 均按 px 语义；中文字体必须同时写 a:ea，否则中文回退宋体"""
    tb = slide.shapes.add_textbox(E(x), E(y), E(w), E(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = ANCH[anchor]
    for i, ln in enumerate(s.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ALIGN[align]
        p.space_before = Pt(0); p.space_after = Pt(0)
        if lh:
            p.line_spacing = lh / 1.2
        r = p.add_run(); r.text = ln
        f = r.font
        f.size = Pt(size * 0.75)
        f.bold = bold
        f.name = fam
        f.color.rgb = RGBColor.from_string(color)
        rPr = r._r.get_or_add_rPr()
        ea = rPr.find(qn('a:ea'))
        if ea is None:
            ea = rPr.makeelement(qn('a:ea'), {}); rPr.append(ea)
        ea.set('typeface', fam)
        if spc:
            rPr.set('spc', str(int(size * 0.75 * spc * 100)))
    return tb


def picture(slide, path, bx, by, bw, bh):
    """等比 contain 到 (bx,by,bw,bh) 并居中 —— 绝不拉伸"""
    with Image.open(path) as im:
        iw, ih = im.size
    r = min(bw / iw, bh / ih)
    w, h = iw * r, ih * r
    x, y = bx + (bw - w) / 2, by + (bh - h) / 2
    slide.shapes.add_picture(path, E(x), E(y), E(w), E(h))
    return x, y, w, h


def header(slide, left, right):
    """页眉：左品牌字 + 右侧英文 + 连接细线（分隔线落在页眉带垂直中心）"""
    size = 23
    lw = tw(left, size, True)
    text(slide, MX, HDR / 2 - size, 900, size * 2, left, size, INK, bold=True,
         anchor='m', spc=0.01)
    rw = tw(right, size, True, FEN)
    rx = W - MX - rw
    text(slide, rx - 2, HDR / 2 - size, rw + 4, size * 2, right, size, INK,
         bold=True, anchor='m', align='r', fam=FEN, spc=0.01)
    rect(slide, MX + lw + 30, HDR / 2 - 0.75, rx - 30 - (MX + lw + 30), 1.5, INK)


def bg(slide, color):
    f = slide.background.fill
    f.solid(); f.fore_color.rgb = RGBColor.from_string(color)


def new(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


# ---------- 文档 ----------
prs = Presentation()
prs.slide_width = E(W)
prs.slide_height = E(H)
PICS = {v['key']: f"logo-{v['key']}.png" for v in VERSIONS}

# ============ P1-P3 · 每版详页 ============
for v in VERSIONS:
    s = new(prs)
    bg(s, PAPER)
    header(s, f'{BRAND_NAME} · 标志释义 {v["no"]}', BRAND_EN)

    # 左：标志展示台
    rect(s, MX, 158, 560, 442, PLATE)
    picture(s, os.path.join(A, PICS[v['key']]), MX + 34, 158 + 34, 560 - 68, 442 - 68)

    # 右：标志构成三段
    rx = MX + 560 + 60.4          # 770.0
    rw = W - MX - rx              # 440.4
    rect(s, rx, 158, 56, 3, BRAND)
    text(s, rx, 176, rw, 42, '标志构成', 25, INK, bold=True, anchor='m')
    yy = 244                       # 段节奏：步进 120，正文 lh 1.62（再紧会粘连）
    for t, d in v['blocks']:
        text(s, rx, yy, rw, 28, t, 18, BRAND_D, bold=True, anchor='m')
        text(s, rx, yy + 34, rw, 80, d, 15, INK2, align='j', lh=1.62)
        yy += 120

    # 下：色彩应用（上排等宽色卡承色值，下排分宽条承用量）
    rect(s, MX, 630, CW, 1, LINEC)
    text(s, MX, 644, 400, 34, '色彩应用', 21, INK, bold=True, anchor='m')
    text(s, MX + CW - 760, 644, 760, 34, COLOR_HINT, 13, INK2, align='r', anchor='m')

    cw = (CW - 3 * 20) / 4
    for i, (nm, hx, pct, tc) in enumerate(v['colors']):
        x = MX + i * (cw + 20)
        rect(s, x, 686, cw, 66, hx)
        frame(s, x, 686, cw, 66)
        text(s, x + 16, 699, cw - 32, 24, nm, 14, tc, bold=True)
        text(s, x + 16, 721, cw - 32, 22, '#' + hx, 12, tc, fam=FMONO)

    acc = MX
    for nm, hx, pct, tc in v['colors']:
        bw = CW * pct / 100.0
        rect(s, acc, 762, bw, 22, hx)
        frame(s, acc, 762, bw, 22)
        if bw > 72:                # 太窄的段放不下百分比，自动略过
            ltc = tc
            text(s, acc, 762, bw - 12, 22, f'{pct}%', 13, ltc, bold=True,
                 align='r', anchor='m', fam=FEN)
        acc += bw

    text(s, MX, 800, CW, 24, v['note'], 13, INK2)

# ============ P4 · 并置总览 ============
s = new(prs)
bg(s, PAPER)
header(s, f'{BRAND_NAME} · 三版并置', BRAND_EN)
rect(s, MX, 158, 56, 3, BRAND)
text(s, MX, 176, 700, 42, LEAD_P4_TITLE, 22, INK, bold=True, anchor='m')
text(s, MX + CW - 620, 176, 620, 42, LEAD_P4, 13.5, INK2, align='r', anchor='m')

n = len(VERSIONS)
gw = (CW - (n - 1) * 36) / n
for i, v in enumerate(VERSIONS):
    x = MX + i * (gw + 36)
    rect(s, x, 236, gw, 414, PLATE)
    picture(s, os.path.join(A, PICS[v['key']]), x + 26, 236 + 26, gw - 52, 414 - 52)
    text(s, x, 668, gw, 24, v['no'], 15, BRAND, bold=True, align='c', anchor='m', spc=0.20)
    text(s, x, 694, gw, 30, v['name'], 19, INK, bold=True, align='c', anchor='m', spc=0.04)
    text(s, x, 726, gw, 24, v['tags'], 13, INK2, align='c', anchor='m')
rect(s, MX, 770, CW, 1, LINEC)
text(s, MX, 786, CW, 24, FOOT_P4, 13, INK2, align='c')

# ============ P5 · 综合应用 ============
# 五格载体，按 VERSIONS 顺序取 A=灯箱反白 / B=碗盖彩 / C=围裙反白，另加 A 墨稿、C 名片双面
s = new(prs)
bg(s, PAPER)
header(s, f'{BRAND_NAME} · 三版综合应用', BRAND_EN)
K = [v['key'] for v in VERSIONS]
kA, kB, kC = (K + K)[:3]

x, y = MX, 150
rect(s, x, y, 620, 280, 'EDE7DA')                      # ① 门头灯箱
rect(s, x + 40, y + 26, 540, 186, DARK)
picture(s, os.path.join(A, f'logo-{kA}-reverse.png'), x + 60, y + 36, 500, 166)
rect(s, x + 214, y + 226, 192, 54, '3A342E')
rect(s, x + 309, y + 226, 2, 54, '575047')
oval(s, x + 96, y + 224, 24, 24, 'F0E4C8')
oval(s, x + 500, y + 224, 24, 24, 'F0E4C8')

x = 793.6
rect(s, x, y, 416.8, 280, 'F3EFE4')                    # ② 外带碗盖
cx, cy = x + 150, y + 140
oval(s, cx - 120, cy - 120, 240, 240, 'FFFFFF', 'D8D2C4', 1.0)
oval(s, cx - 92, cy - 92, 184, 184, 'FBF9F5', 'E4DED2', 0.8)
picture(s, os.path.join(A, f'logo-{kB}.png'), cx - 72, cy - 72, 144, 144)
text(s, x + 302, y + 96, 108, 24, 'φ120mm', 15, INK, bold=True, fam=FEN)
text(s, x + 302, y + 126, 108, 78, '覆膜碗盖\n双色印刷\n彩色主稿', 12, INK2, lh=1.8)

for lx, lw, t1, t2 in [(MX, 620, '门头灯箱', f'{VERSIONS[0]["no"]}版 · 反白稿'),
                       (793.6, 416.8, '外带碗盖', f'{VERSIONS[1]["no"]}版 · 彩色稿')]:
    text(s, lx, 438, lw, 24, t1, 14, INK, bold=True, align='c', anchor='m')
    text(s, lx, 462, lw, 22, t2, 12.5, INK2, align='c', anchor='m')

gw2 = (CW - 2 * 24) / 3
gy = 486
x = MX                                                  # ③ 餐具烙印（墨稿）
rect(s, x, gy, gw2, 268, 'E8E2D5')
rect(s, x + 58, gy + 44, 222, 180, 'D2B48C', 'BFA075', 1.0, radius=0.06)
for i in range(3):
    rect(s, x + 76, gy + 76 + i * 52, 186, 2, 'C4A67C')
picture(s, os.path.join(A, f'logo-{kA}-mono.png'), x + 86, gy + 62, 166, 144)

x = MX + gw2 + 24                                       # ④ 员工围裙（反白）
rect(s, x, gy, gw2, 268, DARK)
rect(s, x + 100, gy + 56, 138, 180, '992D27')
rect(s, x + 112, gy + 28, 16, 32, '992D27')
rect(s, x + 210, gy + 28, 16, 32, '992D27')
picture(s, os.path.join(A, f'logo-{kC}-reverse.png'), x + 122, gy + 99, 94, 94)

x = MX + 2 * (gw2 + 24)                                 # ⑤ 名片（彩稿 / 反白）
rect(s, x, gy, gw2, 268, 'E8E2D5')
cwd2, chd = 150, 94
cyy = gy + (268 - chd) / 2
lx2 = x + (gw2 - 2 * cwd2 - 20) / 2
rect(s, lx2, cyy, cwd2, chd, 'FFFFFF', 'D8D2C4', 1.0)
picture(s, os.path.join(A, f'logo-{kC}.png'), lx2 + 45, cyy + 17, 60, 60)
rx2 = lx2 + cwd2 + 20
rect(s, rx2, cyy, cwd2, chd, BRAND)
picture(s, os.path.join(A, f'logo-{kC}-reverse.png'), rx2 + 45, cyy + 17, 60, 60)

for i, (t1, t2) in enumerate([('餐具烙印', f'{VERSIONS[0]["no"]}版墨稿 · 单色烫印'),
                              ('员工围裙', f'{VERSIONS[2]["no"]}版 · 反白稿'),
                              ('名片', f'{VERSIONS[2]["no"]}版 · 彩色 / 反白双面')]):
    lx = MX + i * (gw2 + 24)
    text(s, lx, 768, gw2, 24, t1, 14, INK, bold=True, align='c', anchor='m')
    text(s, lx, 792, gw2, 22, t2, 12.5, INK2, align='c', anchor='m')

text(s, MX, 826, CW, 22, FOOT_P5, 12.5, INK2, align='c')

prs.save(OUT)
print('saved:', OUT, os.path.getsize(OUT), 'bytes |', len(prs.slides._sldIdLst), 'slides')
print('!! 收尾必做：转 PDF 渲染逐页核对 + 跑 probe-pptx.py 查图片零变形')
