# -*- coding: utf-8 -*-
"""
deck2pptx.py —— VI 全案母版几何 -> 原生可编辑 PPTX 骨架

配套规范：04-知识沉淀/PPTX原生构建规范.md
用法：拷走本文件，改 CFG 与 PAGES 段，跑：
    python deck2pptx.py
验收：LibreOffice 转 PDF -> PyMuPDF 渲染 PNG -> PIL 像素扫描比对设计值

核心换算：1 px = 9525 EMU（96dpi），母版 px 值可 1:1 直落
"""
import os
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import Image, ImageFont

# ============ CFG：每单只改这里 ============
OUT = r'D:\path\to\output\deck.pptx'
BRAND = 'C63021'        # 品牌主色
BRAND_D = '992D27'      # 品牌辅色
INK = '1C1C16'; INK2 = '6E6A5C'
PAPER = 'FFFFFF'; WARM = 'F3EFE4'; DARK = '1D191A'; LINEC = 'C6C2B8'
FCN = 'Microsoft YaHei' # 中文（必设 a:ea）
FEN = 'Arial'           # 英文
FMONO = 'Consolas'      # 等宽
FSER = 'SimSun'         # 衬线
FPATH = {
    (FCN, True): r'C:\Windows\Fonts\msyhbd.ttc',
    (FCN, False): r'C:\Windows\Fonts\msyh.ttc',
    (FEN, False): r'C:\Windows\Fonts\arial.ttf',
    (FEN, True): r'C:\Windows\Fonts\arialbd.ttf',
    (FSER, False): r'C:\Windows\Fonts\simsun.ttc',
}
# ==========================================

PXU = 9525
W, H = 1360, 850                 # 16:10 画布
MX = W * 0.11                    # 149.6  左右边距 11%
HDR = H * 0.15                   # 127.5  页眉带高 15%
CW = W - 2 * MX                  # 1060.8 版心宽

_fc = {}


def _font(fam, bold):
    k = (fam, bold)
    if k not in _fc:
        p = FPATH.get(k) or FPATH[(FCN, bold)]
        try:
            _fc[k] = ImageFont.truetype(p, 100)
        except OSError:
            _fc[k] = ImageFont.truetype(FPATH[(FCN, bold)], 100)
    return _fc[k]


def E(v):
    return Emu(int(round(v * PXU)))


def tw(s, px, bold=False, fam=FCN):
    """文字像素宽度 —— 页眉花线的起止必须按实测宽算，不能估"""
    return _font(fam, bold).getlength(s) * px / 100


# ---------- 基础件 ----------
def rect(slide, x, y, w, h, color):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, E(x), E(y), E(w), E(h))
    s.fill.solid(); s.fill.fore_color.rgb = RGBColor.from_string(color)
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def frame(slide, x, y, w, h, line=LINEC, weight=1.0, fill=None):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, E(x), E(y), E(w), E(h))
    if fill:
        s.fill.solid(); s.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        s.fill.background()
    s.line.color.rgb = RGBColor.from_string(line)
    s.line.width = Pt(weight * 0.75)
    s.shadow.inherit = False
    return s


ALIGN = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT, 'j': PP_ALIGN.JUSTIFY}
ANCH = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}


def text(slide, x, y, w, h, s, size, color, bold=False, align='l', fam=FCN,
         spc=0.0, lh=None, anchor='t', italic=False):
    """size 传 CSS px；内部转 pt(×0.75)；lh 传 CSS line-height（内部 /1.2）"""
    tb = slide.shapes.add_textbox(E(x), E(y), E(w), E(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = ANCH[anchor]
    for i, ln in enumerate(s.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ALIGN[align]
        p.space_before = Pt(0); p.space_after = Pt(0)
        if lh:
            p.line_spacing = lh / 1.2
        r = p.add_run(); r.text = ln
        f = r.font
        f.size = Pt(size * 0.75)
        f.bold = bold; f.italic = italic
        f.name = fam
        f.color.rgb = RGBColor.from_string(color)
        rPr = r._r.get_or_add_rPr()
        ea = rPr.find(qn('a:ea'))          # ← 中文字体槽位，不写就回退宋体
        if ea is None:
            ea = rPr.makeelement(qn('a:ea'), {}); rPr.append(ea)
        ea.set('typeface', fam)
        if spc:                            # 字距，单位 1/100 pt
            rPr.set('spc', str(int(size * 0.75 * spc * 100)))
    return tb


def picture(slide, path, bx, by, bw, bh):
    """等比 contain + 居中，绝不拉伸"""
    with Image.open(path) as im:
        iw, ih = im.size
    r = min(bw / iw, bh / ih)
    w, h = iw * r, ih * r
    slide.shapes.add_picture(path, E(bx + (bw - w) / 2), E(by + (bh - h) / 2), E(w), E(h))
    return w, h


def bg(slide, color):
    f = slide.background.fill
    f.solid(); f.fore_color.rgb = RGBColor.from_string(color)


def header(slide, left, right, dark=False):
    """母版页眉：左标题 + 伸缩线 + 右英文，垂直居中于 0~HDR"""
    col = PAPER if dark else INK
    size = 23
    lw = tw(left, size, True)
    text(slide, MX, HDR / 2 - size, 900, size * 2, left, size, col, bold=True, anchor='m', spc=0.01)
    rw = tw(right, size, True, FEN)
    rx = W - MX - rw
    text(slide, rx - 2, HDR / 2 - size, rw + 4, size * 2, right, size, col,
         bold=True, anchor='m', align='r', fam=FEN, spc=0.01)
    rect(slide, MX + lw + 30, HDR / 2 - 0.75, rx - 30 - (MX + lw + 30), 1.5, col)


def section(slide, en1, en2, pill, side_l='', side_r=''):
    """母版章节隔断：双字圆环 + 中央胶囊"""
    bg(slide, BRAND)
    cx, cy = W / 2, (HDR + H) / 2
    o = slide.shapes.add_shape(MSO_SHAPE.OVAL, E(cx - 146), E(cy - 146), E(292), E(292))
    o.fill.background(); o.line.color.rgb = RGBColor.from_string(PAPER)
    o.line.width = Pt(1.2); o.shadow.inherit = False
    text(slide, cx - 146, cy - 146 + 44, 292, 24, en1, 20, PAPER, align='c', anchor='m', fam=FSER, spc=0.42)
    text(slide, cx - 146, cy + 146 - 68, 292, 24, en2, 20, PAPER, align='c', anchor='m', fam=FSER, spc=0.42)
    pw = tw(pill, 37, True) + 3 * 37 * 0.10 + 60
    ph = 37 * 1.2 + 26
    rect(slide, cx - pw / 2, cy - ph / 2, pw, ph, PAPER)
    text(slide, cx - pw / 2, cy - ph / 2, pw, ph, pill, 37, BRAND, bold=True, align='c', anchor='m', spc=0.10)
    if side_l:
        text(slide, MX, cy - 14, 350, 28, side_l, 20, PAPER, align='r', anchor='m', fam=FSER, italic=True)
    if side_r:
        text(slide, W - MX - 350, cy - 14, 350, 28, side_r, 20, PAPER, align='l', anchor='m', fam=FSER, italic=True)


def new(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])   # 6 = 空白版式


# ================= 主程序 =================
if __name__ == '__main__':
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    prs = Presentation()
    prs.slide_width = E(W)
    prs.slide_height = E(H)

    # --- 示例：封面 ---
    s = new(prs); bg(s, BRAND)
    text(s, 0, 246, W, 112, '品牌名', 88, PAPER, bold=True, align='c', anchor='m', spc=0.10)
    text(s, 0, 370, W, 56, '品牌视觉识别系统', 34, PAPER, bold=True, align='c', anchor='m', spc=0.06)
    text(s, 0, 432, W, 30, 'BRAND VISUAL IDENTITY SYSTEM', 15, 'FFE9E4',
         bold=True, align='c', anchor='m', fam=FEN, spc=0.38)

    # --- 示例：单件居中图页 ---
    s = new(prs); bg(s, PAPER)
    header(s, '品牌名 · 标志主形象', 'BRAND')
    picture(s, r'D:\path\to\logo.png', MX, 154, CW, 620)
    text(s, MX, 796, CW, 24, '图注', 17, INK, bold=True, align='c', anchor='m')
    text(s, MX, 824, CW, 22, '一句话说明', 13.5, INK2, align='c', anchor='m')

    # --- 示例：章节隔断 ---
    s = new(prs)
    section(s, 'MARK', 'SYSTEM', '标志部分')

    # --- 示例：五色卡 ---
    s = new(prs); bg(s, PAPER)
    header(s, '品牌名 · 品牌标准色', 'BRAND')
    SW = [('主色名', 'C63021', '主色 · 用途', PAPER), ('辅色名', '992D27', '辅助 · 用途', PAPER),
          ('墨色', '1C1C16', '正文 / 字标', PAPER), ('底色', 'F3EFE4', '页面底色', INK),
          ('中性', '8A8578', '分隔 / 次级', PAPER)]
    cwd, cy_, chh = CW / 5, 275, 380
    for i, (nm, hx, role, tc) in enumerate(SW):
        x = MX + i * cwd
        rect(s, x, cy_, cwd, chh, hx)
        frame(s, x, cy_, cwd, chh)
        text(s, x + 16, cy_ + chh - 96, cwd - 32, 26, nm, 17, tc, bold=True)
        text(s, x + 16, cy_ + chh - 68, cwd - 32, 22, '#' + hx, 13, tc, fam=FMONO)
        text(s, x + 16, cy_ + chh - 40, cwd - 32, 24, role, 12.5, tc)

    # --- 示例：尾页 ---
    s = new(prs); bg(s, BRAND)
    text(s, MX, 52, 500, 46, 'CHUANG MENG PLANNING\n创梦餐饮品牌策划设计', 13, PAPER,
         bold=True, lh=1.65, spc=0.22)
    text(s, MX, H * 0.5 - 90, 900, 180, 'Thanks!', 112, PAPER, fam=FSER, anchor='m')

    prs.save(OUT)
    print('saved:', OUT, os.path.getsize(OUT), 'bytes')
